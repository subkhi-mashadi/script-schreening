"""Pembuat dokumen uji secara programatik (tanpa berkas biner di repo)."""

import io
import zipfile

import docx
import pymupdf
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

PARA = (
    "Penelitian ini membahas pengelolaan naskah buku secara sistematis di lingkungan penerbit. "
    "Setiap tahap pemeriksaan dicatat agar keputusan editorial dapat ditelusuri kembali. "
)


def make_docx(*, footnote: bool = False, tracked: bool = False, pattern_heading: bool = False, title: str = "Judul Uji") -> bytes:
    d = docx.Document()
    d.core_properties.title = title
    d.core_properties.author = "Penulis Uji"
    d.add_paragraph(title, style="Title")
    d.add_heading("KATA PENGANTAR", level=1)
    d.add_paragraph(PARA * 2)
    if pattern_heading:
        p = d.add_paragraph()
        p.add_run("BAB I").bold = True
        p2 = d.add_paragraph()
        p2.add_run("PENDAHULUAN").bold = True
    else:
        d.add_heading("BAB I", level=1)
        d.add_heading("PENDAHULUAN", level=1)
    d.add_heading("1.1 Latar Belakang", level=2)
    body = d.add_paragraph(PARA * 3)
    d.add_paragraph("Butir pertama daftar", style="List Bullet")
    d.add_paragraph("Butir kedua daftar", style="List Bullet")
    d.add_paragraph("Kutipan panjang dari sumber.", style="Quote")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Kolom A"
    t.cell(0, 1).text = "Kolom B"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "2"
    d.add_heading("1.2 Rumusan Masalah", level=2)
    d.add_paragraph(PARA * 2)
    d.add_page_break()
    d.add_heading("BAB II KAJIAN PUSTAKA", level=1)
    d.add_paragraph(PARA * 4)
    d.add_heading("DAFTAR PUSTAKA", level=1)
    d.add_paragraph("Sutrisno, A. (2020). Manajemen Penerbitan. Jakarta: Penerbit Contoh.")

    if tracked:
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), "1")
        ins.set(qn("w:author"), "Editor")
        r = OxmlElement("w:r")
        t_el = OxmlElement("w:t")
        t_el.text = " TEKS SISIPAN"
        r.append(t_el)
        ins.append(r)
        body._p.append(ins)
        dele = OxmlElement("w:del")
        dele.set(qn("w:id"), "2")
        dele.set(qn("w:author"), "Editor")
        r2 = OxmlElement("w:r")
        dt = OxmlElement("w:delText")
        dt.text = " TEKS TERHAPUS"
        r2.append(dt)
        dele.append(r2)
        body._p.append(dele)

    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    if footnote:
        data = _inject_footnote(data)
    return data


def _inject_footnote(data: bytes) -> bytes:
    """Tambahkan footnotes.xml + referensi pada paragraf isi pertama setelah 'Latar Belakang'."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    footnotes_xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{w}">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>Catatan kaki uji tentang sumber data.</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "word/document.xml":
                text = content.decode()
                marker = "Penelitian ini membahas"
                idx = text.find(marker, text.find("Latar Belakang"))
                run_end = text.find("</w:r>", idx) + len("</w:r>")
                ref = '<w:r><w:footnoteReference w:id="1"/></w:r>'
                content = (text[:run_end] + ref + text[run_end:]).encode()
            elif item.filename == "word/_rels/document.xml.rels":
                text = content.decode()
                rel = ('<Relationship Id="rIdFn1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                       'relationships/footnotes" Target="footnotes.xml"/>')
                content = text.replace("</Relationships>", rel + "</Relationships>").encode()
            elif item.filename == "[Content_Types].xml":
                text = content.decode()
                ov = ('<Override PartName="/word/footnotes.xml" ContentType="application/'
                      'vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>')
                content = text.replace("</Types>", ov + "</Types>").encode()
            dst.writestr(item, content)
        dst.writestr("word/footnotes.xml", footnotes_xml)
    return out.getvalue()


def make_pdf(*, with_toc: bool = True, pages: int = 4) -> bytes:
    doc = pymupdf.open()
    toc = []
    for i in range(pages):
        page = doc.new_page()
        page.insert_text((72, 30), "Judul Buku Uji — Header Berjalan", fontsize=8)
        y = 90
        if i % 2 == 0:
            title = f"BAB {['I', 'II', 'III', 'IV', 'V'][i // 2]} BAGIAN {i // 2 + 1}"
            page.insert_text((72, y), title, fontsize=18, fontname="hebo")
            toc.append([1, title, i + 1])
            y += 40
        page.insert_text((72, y), f"{i // 2 + 1}.1 Subbagian Halaman {i + 1}", fontsize=14, fontname="hebo")
        if i % 2 == 0:
            toc.append([2, f"{i // 2 + 1}.1 Subbagian Halaman {i + 1}", i + 1])
        y += 30
        rect = pymupdf.Rect(72, y, 520, y + 300)
        page.insert_textbox(rect, PARA * 4, fontsize=11)
        page.insert_text((300, 800), str(i + 1), fontsize=9)
    if with_toc:
        doc.set_toc(toc)
    data = doc.tobytes()
    doc.close()
    return data


def make_scanned_pdf(pages: int = 3) -> bytes:
    doc = pymupdf.open()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), False)
    pix.set_rect(pix.irect, (200, 200, 200))
    png = pix.tobytes("png")
    for _ in range(pages):
        page = doc.new_page()
        page.insert_image(page.rect, stream=png)
    data = doc.tobytes()
    doc.close()
    return data


def make_txt() -> bytes:
    text = (
        "KATA PENGANTAR\n\n" + PARA * 2 + "\n\n"
        "BAB I\n\nPENDAHULUAN\n\n" + PARA * 3 + "\n\n"
        "1.1 Latar Belakang\n\n" + PARA * 3 + "\n\n"
        "DAFTAR PUSTAKA\n\nSutrisno, A. (2020). Manajemen Penerbitan.\n"
    )
    return text.encode("utf-8")


def make_zip_bomb_docx() -> bytes:
    base = make_docx()
    src = zipfile.ZipFile(io.BytesIO(base))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("word/media/padding.bin", b"\0" * (20 * 1024 * 1024))
    return out.getvalue()


def make_macro_docx() -> bytes:
    base = make_docx()
    src = zipfile.ZipFile(io.BytesIO(base))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("word/vbaProject.bin", b"macro")
    return out.getvalue()


__all__ = ["Pt", "make_docx", "make_macro_docx", "make_pdf", "make_scanned_pdf", "make_txt", "make_zip_bomb_docx"]


STORY = (
    "Pagi itu Kona berjalan menyusuri pematang sawah sambil membawa keranjang kecil. Angin bertiup pelan dan "
    "burung-burung beterbangan di atas pohon kelapa. Ia memikirkan janji yang dibuat bersama adiknya kemarin sore. "
)


def make_book_docx(
    *,
    issues: bool = False,
    scientific: bool = False,
    chapters: int = 3,
    title: str = "Kisah Pondok Kecil",
    author: str = "Rina Pratiwi",
) -> bytes:
    """Naskah lengkap: halaman judul, halaman hak cipta, kata pengantar, daftar isi, bab, biodata.

    `issues=True` menambahkan masalah yang diketahui untuk diuji: placeholder, duplikasi,
    kata tidak baku, kata depan tergabung, spasi ganda, bab melompat, entri daftar isi terpecah.
    """
    d = docx.Document()
    d.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
    d.core_properties.title = title
    d.add_paragraph(title)
    d.add_paragraph(f"Penulis: {author}")
    d.add_paragraph("Penerbit Dimensi Antara Pustaka")
    if issues:  # halaman judul tercetak dua kali
        d.add_paragraph(title)
        d.add_paragraph(f"Penulis: {author}")
        d.add_paragraph("Penerbit Dimensi Antara Pustaka")
    d.add_paragraph("Hak cipta © 2026 pada penulis. Diterbitkan oleh PT Dimensi Antara Pustaka.")
    d.add_heading("KATA PENGANTAR", level=1)
    d.add_paragraph(STORY)
    d.add_heading("DAFTAR ISI", level=1)
    numbers = [1, 2, 4] if issues else list(range(1, chapters + 1))
    for n in numbers:
        if issues and n == 2:
            d.add_paragraph(f"BAB {n}\t{n * 10}", style="TOC 1")
            d.add_paragraph(f"JUDUL BAB {n}\t{n * 10}", style="TOC 1")
        else:
            d.add_paragraph(f"BAB {n} JUDUL BAB {n}\t{n * 10}", style="TOC 1")
    for n in numbers:
        d.add_heading(f"BAB {n} JUDUL BAB {n}", level=1)
        if scientific and n == 1:
            d.add_heading("Pendahuluan", level=2)
        d.add_heading(f"{n}.1 Bagian Pertama", level=2)
        text = STORY * 6
        if scientific:
            text += " Hal ini sejalan dengan temuan sebelumnya (Sutrisno, 2020) dan diperkuat oleh Lestari (2019)."
            if n == 2:
                text += " Pendapat lain dikemukakan (Hidayat, 2018)."
        d.add_paragraph(text)
        if issues and n == 1:
            d.add_paragraph("TODO: lengkapi bagian ini.  Aktifitas warga disana sangat ramai.")
            d.add_paragraph(STORY * 3)
        if issues and n == 4:
            d.add_paragraph(STORY * 3)  # paragraf identik dengan bab 1
    if scientific:
        d.add_heading("PENUTUP", level=1)
        d.add_paragraph(STORY * 2)
        d.add_heading("DAFTAR PUSTAKA", level=1)
        d.add_paragraph("Lestari, B. (2019). Kajian Penerbitan. Bandung: Penerbit Contoh.")
        d.add_paragraph("Sutrisno, A. (2020). Manajemen Penerbitan. Jakarta: Penerbit Contoh.")
        d.add_paragraph("Wibowo, C. (2017). Buku yang tidak disitir. Surabaya: Penerbit Contoh.")
    d.add_heading("BIODATA PENULIS", level=1)
    d.add_paragraph(f"{author} lahir di Medan dan menulis cerita anak sejak 2015.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


DIALOG = [
    "“Kak, kita main ke ladang yuk!” ajak Nagin.",
    "“Nanti dulu, Kakak masih belajar.”",
    "“Ayolah, sebentar saja!”",
    "Kona tertawa melihat adiknya.",
    "“Baiklah, tapi pulang sebelum sore.”",
]


def make_story_docx(*, child: bool = False, with_toc: bool = True, with_preface: bool = True, translation: bool = False,
                    title: str = "Petualangan di Ladang") -> bytes:
    """Naskah cerita (banyak dialog). `child` membuat huruf isi 14 pt dan kalimat pendek."""
    d = docx.Document()
    if child:
        d.styles["Normal"].font.size = Pt(14)
    d.add_paragraph(title)
    d.add_paragraph("Penulis: Rina Pratiwi")
    d.add_paragraph("Penerbit Dimensi Antara Pustaka")
    if translation:
        d.add_paragraph("Judul asli: Adventures in the Field. Diterjemahkan dari bahasa Inggris oleh Budi.")
    d.add_paragraph("Hak cipta © 2026 pada penulis. Diterbitkan oleh PT Dimensi Antara Pustaka.")
    if with_preface:
        d.add_heading("KATA PENGANTAR", level=1)
        d.add_paragraph("Cerita ini ditulis untuk menemani waktu membaca bersama keluarga di rumah.")
    if with_toc:
        d.add_heading("DAFTAR ISI", level=1)
        for n in (1, 2):
            d.add_paragraph(f"BAB {n} CERITA {n}\t{n * 5}")
    for n in (1, 2):
        d.add_heading(f"BAB {n} CERITA {n}", level=1)
        for _ in range(6):
            for line in DIALOG:
                d.add_paragraph(line)
            if not child:
                d.add_paragraph(STORY)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
