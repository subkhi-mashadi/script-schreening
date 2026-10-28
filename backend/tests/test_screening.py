import pytest
from sqlalchemy import select, text

from app.core.db import SessionLocal
from app.models import RuleSet
from app.screening.scoring import RuleInput, score
from tests.factories import make_book_docx, make_scanned_pdf
from tests.test_api import API, DOCX_MIME, create_ms, events, upload


def quick(client, headers, data, name="buku.docx", mime=DOCX_MIME, manuscript_type="fiksi", distribution="dijual_umum", **form):
    form = {"manuscript_type": manuscript_type, **({"distribution": distribution} if distribution else {}), **form}
    r = client.post(f"{API}/quick-checks", files={"file": (name, data, mime)}, data=form, headers=headers)
    assert r.status_code == 201, r.text
    return r.json()


def latest_run(client, headers, version_id):
    r = client.get(f"{API}/versions/{version_id}/screenings/latest", headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def layout_run(client, headers, version_id):
    """Tahap 2: cek layout atas permintaan."""
    r = client.post(f"{API}/versions/{version_id}/layout-check", headers=headers)
    assert r.status_code == 202, r.text
    run = client.get(f"{API}/versions/{version_id}/screenings/latest", params={"scope": "layout"}, headers=headers).json()
    assert run["scope"] == "layout"
    return run


def findings(client, headers, run_id, **params):
    return client.get(f"{API}/screenings/{run_id}/findings", params={"page_size": 500, **params}, headers=headers).json()["items"]


def codes(items):
    return {f["rule_code"] for f in items}


def set_authors(client, headers, ms_id, names):
    r = client.patch(f"{API}/manuscripts/{ms_id}", json={"authors": [{"author_name": n} for n in names]}, headers=headers)
    assert r.status_code == 200, r.text


# --- Alur otomatis ----------------------------------------------------------


def test_upload_runs_isbn_stage_only(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_book_docx())
    run = latest_run(client, op, qc["version"]["id"])
    assert run["status"] == "COMPLETED" and run["scope"] == "isbn"
    assert run["rubric"] is None and run["total_score"] is None and run["editorial_status"] is None
    # Naskah lengkap: semua syarat wajib Pasal 10 ayat (3) huruf b & c terpenuhi → layak.
    assert run["isbn_status"] == "PASS"
    comp = client.get(f"{API}/screenings/{run['id']}/compliance", headers=op).json()
    by = {c["rule_code"]: c for c in comp}
    for code in ("ISB-001", "ISB-002", "ISB-003", "ISB-004", "ISB-010"):
        assert by[code]["observed"] == "pass" and by[code]["result_status"] == "pass" and by[code]["blocking"], code
    assert by["ISB-006"]["observed"] == "unknown" and by["ISB-006"]["requires_human_verification"]
    assert by["ISB-007"]["result_status"] == "not_applicable"
    assert "Pasal 10 ayat (3)" in by["ISB-001"]["source_reference"] and "peraturan.go.id" in by["ISB-001"]["source_url"]
    # Belum ada cek layout sampai diminta.
    assert client.get(f"{API}/versions/{qc['version']['id']}/screenings/latest", params={"scope": "layout"}, headers=op).json() is None
    assert all(f["track"] == "compliance" for f in findings(client, op, run["id"]))
    assert {e.event_type for e in events("screening.queued", "screening.completed")} == {"screening.queued", "screening.completed"}


def test_layout_check_on_request(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_book_docx())
    run = layout_run(client, op, qc["version"]["id"])
    assert run["status"] == "COMPLETED" and run["isbn_status"] is None and run["rubric"]["name"] == "Rubrik fiksi"
    lay = run["result_summary_json"]["layout"]
    assert lay["page_sizes"][0]["width_cm"] == 21.6 and lay["body_fonts"][0]["value"]
    items = findings(client, op, run["id"])
    assert all(f["track"] == "editorial" for f in items)
    # Naskah contoh tanpa nomor halaman → temuan tata letak.
    assert any(f["rule_code"] == "LAY-007" for f in items)
    crit = {s["criterion_code_snapshot"]: s for s in client.get(f"{API}/screenings/{run['id']}/scores", headers=op).json()}
    assert crit["layout_typography"]["scoring_status"] == "scored" and float(crit["layout_typography"]["raw_score"]) <= 80


def test_known_issues_are_found_with_locations(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_book_docx(issues=True))
    isbn = latest_run(client, op, qc["version"]["id"])
    assert isbn["isbn_status"] == "FAIL", "teks TODO berarti naskah belum final (Pasal 10 ayat (3) huruf c)"
    run = layout_run(client, op, qc["version"]["id"])
    items = findings(client, op, run["id"])
    titles = [f["title"] for f in items]
    assert {"TEC-001", "DUP-001", "LAN-001", "LAN-005", "LAN-003", "STR-002", "STR-004"} <= codes(items), titles
    todo = next(f for f in items if f["rule_code"] == "TEC-001")
    assert "TODO" in todo["evidence_excerpt"] and todo["paragraph_id"] and todo["section_path"].startswith("BAB 1 JUDUL BAB 1")
    assert todo["char_end"] > todo["char_start"]
    dup_titles = [f["title"] for f in items if f["rule_code"] == "DUP-001"]
    assert "3 paragraf berurutan terulang" in dup_titles and "Paragraf identik terulang" in dup_titles
    assert any("BAB 4" in f["title"] for f in items if f["rule_code"] == "STR-002")
    assert any("terpecah" in f["title"] for f in items if f["rule_code"] == "STR-004")
    words = {f["machine_result_json"]["details"].get("word") for f in items if f["rule_code"] in ("LAN-005", "LAN-003")}
    assert {"Aktifitas", "disana"} <= words
    assert float(run["total_score"]) < 85 and run["editorial_status"] in ("MINOR_REVISION", "MAJOR_REVISION")
    # Filter peta revisi per bab & keparahan.
    bab1 = findings(client, op, run["id"], section_path="BAB 1 JUDUL BAB 1")
    assert bab1 and all(f["section_path"].startswith("BAB 1 JUDUL BAB 1") for f in bab1)
    facets = client.get(f"{API}/screenings/{run['id']}/findings/facets", headers=op).json()
    assert facets["severity"]["major"] >= 2 and facets["chapter"][0]["count"] >= 1


def test_scores_explain_each_criterion(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_book_docx(issues=True))
    run = layout_run(client, op, qc["version"]["id"])
    rows = client.get(f"{API}/screenings/{run['id']}/scores", headers=op).json()
    total = next(r for r in rows if r["result_type"] == "total")
    crit = [r for r in rows if r["result_type"] == "criterion"]
    assert sum(float(r["weight_percent"]) for r in crit) == 100
    fmt = next(r for r in crit if r["criterion_code_snapshot"] == "format_technical")
    rules = {d["rule_code"]: d for d in fmt["calculation_details_json"]["rules"]}
    assert rules["TEC-001"]["deduction"] == 10 and float(fmt["raw_score"]) <= 90
    recomputed = sum(float(r["weight_percent"]) * float(r["raw_score"]) for r in crit) / 100
    assert abs(recomputed - float(total["raw_score"])) < 0.02
    assert "kriteria dapat dinilai" in total["explanation"]


def test_rerun_is_deterministic_and_keeps_history(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_book_docx(issues=True))
    first = latest_run(client, op, qc["version"]["id"])
    r = client.post(f"{API}/screenings/{first['id']}/rerun", headers=op)
    assert r.status_code == 202
    second = latest_run(client, op, qc["version"]["id"])
    assert second["run_number"] == 2 and second["id"] != first["id"]
    strip = lambda xs: sorted((f["rule_code"], f["title"], f["paragraph_id"], f["char_start"], f["severity"]) for f in xs)  # noqa: E731
    assert strip(findings(client, op, first["id"])) == strip(findings(client, op, second["id"]))
    assert first["total_score"] == second["total_score"] and first["editorial_status"] == second["editorial_status"]
    runs = client.get(f"{API}/manuscripts/{qc['manuscript_id']}/screenings", headers=op).json()
    assert [x["run_number"] for x in runs] == [2, 1]


def test_scanned_pdf_screening_is_incomplete_not_scored(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_scanned_pdf(), name="pindai.pdf", mime="application/pdf")
    run = latest_run(client, op, qc["version"]["id"])
    assert run["status"] == "INCOMPLETE" and run["isbn_status"] == "INCOMPLETE" and run["editorial_status"] is None
    assert findings(client, op, run["id"]) == []
    lay = layout_run(client, op, qc["version"]["id"])
    assert lay["editorial_status"] == "NOT_SCORABLE" and lay["total_score"] is None


def test_scientific_citations_checked_and_not_for_fiction(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_book_docx(scientific=True), manuscript_type="buku_ilmiah")
    run = layout_run(client, op, qc["version"]["id"])
    items = findings(client, op, run["id"])
    missing = [f for f in items if f["rule_code"] == "REF-002"]
    assert [f["machine_result_json"]["details"]["name"] for f in missing] == ["Hidayat"]
    assert missing[0]["requires_human_verification"] is True
    uncited = [f for f in items if f["rule_code"] == "REF-003"]
    assert len(uncited) == 1 and "Wibowo" in uncited[0]["evidence_excerpt"]
    assert run["rubric"]["name"] == "Rubrik buku ilmiah/ajar"
    fiction = layout_run(client, op, quick(client, op, make_book_docx(scientific=True))["version"]["id"])
    assert not {"REF-002", "REF-003"} & codes(findings(client, op, fiction["id"]))


def test_isbn_verdict_rules(client, auth):
    import io

    import docx as _docx

    op = auth("intake_operator")
    assert latest_run(client, op, quick(client, op, make_book_docx())["version"]["id"])["isbn_status"] == "PASS"

    def without(text: str) -> bytes:
        d = _docx.Document(io.BytesIO(make_book_docx()))
        for p in d.paragraphs:
            if text in p.text:
                p._p.getparent().remove(p._p)
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    # Tanpa halaman hak cipta (Pasal 10 ayat (3) huruf b angka 2) → belum layak.
    run = latest_run(client, op, quick(client, op, without("Hak cipta"))["version"]["id"])
    comp = {c["rule_code"]: c for c in client.get(f"{API}/screenings/{run['id']}/compliance", headers=op).json()}
    assert run["isbn_status"] == "FAIL" and comp["ISB-002"]["result_status"] == "fail"
    # Tanpa kata pengantar → belum layak untuk nonfiksi; fiksi dikecualikan (Juknis 4.3.3).
    assert latest_run(client, op, quick(client, op, without("KATA PENGANTAR"), manuscript_type="nonfiksi_umum")["version"]["id"])["isbn_status"] == "FAIL"
    assert latest_run(client, op, quick(client, op, without("KATA PENGANTAR"), manuscript_type="fiksi")["version"]["id"])["isbn_status"] == "PASS"

    # Peringatan non-penentu (mis. jenis terbitan skripsi) tidak menggagalkan kelayakan.
    d = _docx.Document(io.BytesIO(make_book_docx()))
    d.paragraphs[0].insert_paragraph_before("SKRIPSI")
    buf = io.BytesIO()
    d.save(buf)
    run = latest_run(client, op, quick(client, op, buf.getvalue())["version"]["id"])
    comp = {c["rule_code"]: c for c in client.get(f"{API}/screenings/{run['id']}/compliance", headers=op).json()}
    assert comp["ISB-011"]["observed"] == "fail" and not comp["ISB-011"]["blocking"] and run["isbn_status"] == "PASS"

    # Bila verifikasi aturan wajib dicabut admin, status kembali "perlu verifikasi".
    admin = auth("rule_admin")
    rule = next(r for r in client.get(f"{API}/rules", headers=admin).json() if r["rule_code"] == "ISB-001")
    client.post(f"{API}/rules/{rule['id']}/verification", json={"verification_status": "unverified", "note": "ditinjau ulang"}, headers=admin)
    rs = client.post(f"{API}/rule-sets", json={"change_notes": "ISB-001 ditinjau ulang"}, headers=admin).json()
    client.post(f"{API}/rule-sets/{rs['id']}/simulate", files={"file": ("a.docx", make_book_docx(), DOCX_MIME)}, data={"manuscript_type": "fiksi"},
                headers=admin)
    assert client.post(f"{API}/rule-sets/{rs['id']}/publish", headers=admin).status_code == 200
    assert latest_run(client, op, quick(client, op, make_book_docx())["version"]["id"])["isbn_status"] == "NEEDS_VERIFICATION"


# --- Versi aturan & rubrik ----------------------------------------------------


def test_published_rule_set_and_rubric_are_immutable():
    with SessionLocal() as db:
        with pytest.raises(Exception, match="dipublikasikan"):
            db.execute(text("UPDATE rule_sets SET change_notes = 'x' WHERE status = 'published'"))
        db.rollback()
        with pytest.raises(Exception, match="dipublikasikan"):
            db.execute(text("UPDATE rule_set_items SET enabled = false"))
        db.rollback()
        with pytest.raises(Exception, match="dipublikasikan"):
            db.execute(text("UPDATE rubric_criteria SET weight_percent = 50"))
        db.rollback()
        with pytest.raises(Exception, match="dihapus"):
            db.execute(text("DELETE FROM rubric_versions"))
        db.rollback()


def test_rule_change_requires_new_version_and_simulation(client, auth):
    admin = auth("rule_admin")
    op = auth("intake_operator")
    before = layout_run(client, op, quick(client, op, make_book_docx(issues=True))["version"]["id"])
    rule = next(r for r in client.get(f"{API}/rules", headers=admin).json() if r["rule_code"] == "TEC-001")
    client.patch(f"{API}/rules/{rule['id']}", json={"scoring_effect": {"criterion_code": "format_technical", "deduction_per_finding": 30, "max_deduction": 30}},
                 headers=admin)
    # Belum ada rule set baru → screening baru tetap memakai snapshot lama.
    again = client.post(f"{API}/screenings/{before['id']}/rerun", headers=op).json()
    again = client.get(f"{API}/screenings/{again['id']}", headers=op).json()
    assert again["total_score"] == before["total_score"]

    rs = client.post(f"{API}/rule-sets", json={"change_notes": "Potongan placeholder dinaikkan"}, headers=admin).json()
    assert client.post(f"{API}/rule-sets", json={"change_notes": "draft kedua"}, headers=admin).status_code == 409
    r = client.post(f"{API}/rule-sets/{rs['id']}/publish", headers=admin)
    assert r.status_code == 422 and "simulasi" in str(r.json()["error"]["details"]).lower()
    client.post(f"{API}/rule-sets/{rs['id']}/simulate", files={"file": ("a.docx", make_book_docx(issues=True), DOCX_MIME)},
                data={"manuscript_type": "fiksi"}, headers=admin)
    detail = client.get(f"{API}/rule-sets/{rs['id']}", headers=admin).json()
    assert detail["diff"]["changed"][0]["rule_code"] == "TEC-001" and "scoring_effect" in detail["diff"]["changed"][0]["fields"]
    assert client.post(f"{API}/rule-sets/{rs['id']}/publish", headers=admin).status_code == 200
    with SessionLocal() as db:
        assert [x.status for x in db.scalars(select(RuleSet).order_by(RuleSet.version))] == ["archived", "published"]

    after = client.post(f"{API}/screenings/{before['id']}/rerun", headers=op).json()
    after = client.get(f"{API}/screenings/{after['id']}", headers=op).json()
    assert after["rule_set"]["version"] == 2 and float(after["total_score"]) < float(before["total_score"])
    # Hasil lama tidak berubah.
    assert client.get(f"{API}/screenings/{before['id']}", headers=op).json()["total_score"] == before["total_score"]
    assert events("rule.updated") and events("rule_set.published")


def test_rubric_weights_must_total_100(client, auth):
    admin = auth("rule_admin")
    fiksi = client.get(f"{API}/rubrics", params={"manuscript_type": "fiksi"}, headers=admin).json()[0]
    draft = client.post(f"{API}/rubrics/{fiksi['id']}/new-version", headers=admin).json()
    assert draft["version"] == 2 and draft["status"] == "draft"
    client.patch(f"{API}/rubrics/{draft['id']}", json={"criteria": [{"criterion_code": "language_terms", "weight_percent": 30}]}, headers=admin)
    r = client.post(f"{API}/rubrics/{draft['id']}/publish", headers=admin)
    assert r.status_code == 422 and "110" in str(r.json()["error"]["details"])
    client.patch(f"{API}/rubrics/{draft['id']}", json={"criteria": [{"criterion_code": "format_technical", "weight_percent": 5}]}, headers=admin)
    assert client.post(f"{API}/rubrics/{draft['id']}/publish", headers=admin).status_code == 200
    assert client.patch(f"{API}/rubrics/{draft['id']}", json={"min_scoring_coverage": 50}, headers=admin).status_code == 409


def test_official_rule_needs_source_reference_to_verify(client, auth):
    admin = auth("rule_admin")
    rule = next(r for r in client.get(f"{API}/rules", headers=admin).json() if r["rule_code"] == "ISB-001")
    client.patch(f"{API}/rules/{rule['id']}", json={"source_reference": None}, headers=admin)
    r = client.post(f"{API}/rules/{rule['id']}/verification", json={"verification_status": "verified", "note": "cek sumber"}, headers=admin)
    assert r.status_code == 422 and r.json()["error"]["code"] == "SOURCE_REQUIRED"


def test_permissions(client, auth):
    op = auth("intake_operator")
    qc = quick(client, op, make_book_docx())
    run = latest_run(client, op, qc["version"]["id"])
    assert client.get(f"{API}/rules", headers=auth("editor")).status_code == 200
    assert client.post(f"{API}/rule-sets", json={"change_notes": "coba-coba"}, headers=auth("editor")).status_code == 403
    assert client.get(f"{API}/screenings/{run['id']}", headers=auth("rule_admin")).status_code == 404
    assert client.get(f"{API}/screenings/{run['id']}", headers=auth("editor")).status_code == 404  # tidak ditugaskan
    assert client.post(f"{API}/screenings", json={"document_version_id": qc["version"]["id"]}, headers=auth("auditor")).status_code == 403
    # Tidak ada endpoint pengubah skor (FR-07).
    from app.main import app

    for route in app.routes:
        if "score" in getattr(route, "path", ""):
            assert getattr(route, "methods", {"GET"}) <= {"GET", "HEAD"}, route.path


def test_screening_blocked_while_extracting_or_failed(client, auth):
    op = auth("intake_operator")
    ms = create_ms(client, op)
    v = upload(client, op, ms["id"], b"%PDF-1.7 rusak", name="r.pdf", mime="application/pdf").json()
    r = client.post(f"{API}/screenings", json={"document_version_id": v["id"]}, headers=op)
    assert r.status_code == 409 and r.json()["error"]["code"] == "EXTRACTION_FAILED"


# --- Skor (fungsi murni) ------------------------------------------------------

CRIT = [{"code": "a", "name": "A", "weight": 60}, {"code": "b", "name": "B", "weight": 40}]
TH = {"pass": 85, "minor_revision": 70, "major_revision": 55, "exclude_not_applicable": True}


def ri(code, crit, n=0, per=10, cap=30, applicable=True, error=False, sev="minor", vs="UNREVIEWED"):
    return RuleInput(code, code, crit, per, cap, applicable, error,
                     [{"severity": sev, "verification_status": vs, "track": "editorial"} for _ in range(n)])


def test_score_formula_and_caps():
    out = score(CRIT, [ri("r1", "a", n=5), ri("r2", "b", n=1)], TH, 80, True, [])
    a, b = out.criteria
    assert a.score == 70 and b.score == 90  # potongan a dibatasi 30
    assert out.total == round((60 * 70 + 40 * 90) / 100, 2) == 78.0
    assert out.editorial_status == "MINOR_REVISION" and out.coverage == 100


def test_not_applicable_is_excluded_not_full_marks():
    out = score(CRIT, [ri("r1", "a", n=0), ri("r2", "b", applicable=False)], TH, 50, True, [])
    assert out.criteria[1].status == "not_applicable" and out.criteria[1].score is None
    assert out.total == 100 and out.scored_count == 1 and out.applicable_count == 1


def test_low_coverage_withholds_total():
    out = score(CRIT, [ri("r1", "a"), ri("r2", "b", error=True)], TH, 80, True, [])
    assert out.total is None and out.editorial_status == "NEEDS_VERIFICATION" and out.coverage == 60


def test_false_positive_and_info_not_counted():
    out = score(CRIT, [ri("r1", "a", n=2, vs="FALSE_POSITIVE"), ri("r2", "b", n=3, sev="info")], TH, 80, True, [])
    assert out.total == 100


def test_open_critical_blocks_pass():
    crit_finding = [{"severity": "critical", "verification_status": "UNREVIEWED", "track": "editorial"}]
    assert score(CRIT, [ri("r1", "a"), ri("r2", "b")], TH, 80, True, crit_finding).editorial_status == "MAJOR_REVISION"
    major = [{"severity": "major", "verification_status": "UNREVIEWED", "track": "editorial"}]
    assert score(CRIT, [ri("r1", "a"), ri("r2", "b")], TH, 80, True, major).editorial_status == "MINOR_REVISION"
    # Temuan kepatuhan tidak memengaruhi rekomendasi editorial (gerbang terpisah).
    isbn = [{"severity": "critical", "verification_status": "UNREVIEWED", "track": "compliance"}]
    assert score(CRIT, [ri("r1", "a"), ri("r2", "b")], TH, 80, True, isbn).editorial_status == "PASS_RECOMMENDED"


def test_extraction_failure_is_not_scorable_not_zero():
    out = score(CRIT, [], TH, 80, False, [])
    assert out.editorial_status == "NOT_SCORABLE" and out.total is None


# --- Klasifikasi jenis terbitan (tahap 1) ---------------------------------------


def classification(client, headers, run_id):
    return client.get(f"{API}/screenings/{run_id}", headers=headers).json()["result_summary_json"]["classification"]


def test_classifies_children_fiction_and_sets_type(client, auth):
    from tests.factories import make_story_docx

    op = auth("intake_operator")
    qc = quick(client, op, make_story_docx(child=True), manuscript_type="lainnya")
    run = latest_run(client, op, qc["version"]["id"])
    c = classification(client, op, run["id"])
    assert c["jenis_terbitan"]["value"] == "buku_anak" and c["jenis_pustaka"]["value"] == "Fiksi"
    assert c["kelompok_pembaca"]["value"] == "Anak" and c["kategori_buku"]["value"] == "Lepas" and c["applied"] is True
    assert "Juknis" in c["source"]
    ms = client.get(f"{API}/manuscripts/{qc['manuscript_id']}", headers=op).json()
    assert ms["manuscript_type"] == "buku_anak" and ms["metadata_json"]["type_source"] == "auto"
    assert events("manuscript.auto_classified")


def test_fiction_without_preface_and_toc_is_eligible(client, auth):
    from tests.factories import make_story_docx

    op = auth("intake_operator")
    qc = quick(client, op, make_story_docx(with_toc=False, with_preface=False), manuscript_type="lainnya")
    run = latest_run(client, op, qc["version"]["id"])
    comp = {c["rule_code"]: c for c in client.get(f"{API}/screenings/{run['id']}/compliance", headers=op).json()}
    assert classification(client, op, run["id"])["jenis_pustaka"]["value"] == "Fiksi"
    assert comp["ISB-003"]["result_status"] == "not_applicable" and comp["ISB-004"]["result_status"] == "not_applicable"
    assert "Juknis" in comp["ISB-004"]["explanation"]
    assert run["isbn_status"] == "PASS"


def test_nonfiction_without_toc_is_not_eligible(client, auth):
    import io

    import docx as _docx

    op = auth("intake_operator")
    d = _docx.Document(io.BytesIO(make_book_docx()))
    for p in d.paragraphs:
        if p.text == "DAFTAR ISI" or p.style.name.startswith("TOC"):
            p._p.getparent().remove(p._p)
    buf = io.BytesIO()
    d.save(buf)
    run = latest_run(client, op, quick(client, op, buf.getvalue(), manuscript_type="nonfiksi_umum")["version"]["id"])
    assert run["isbn_status"] == "FAIL"


def test_translation_detected_requires_permission_document(client, auth):
    from tests.factories import make_story_docx

    op = auth("intake_operator")
    qc = quick(client, op, make_story_docx(translation=True), manuscript_type="lainnya")
    run = latest_run(client, op, qc["version"]["id"])
    assert classification(client, op, run["id"])["terjemahan"]["value"] == "Ya"
    comp = {c["rule_code"]: c for c in client.get(f"{API}/screenings/{run['id']}/compliance", headers=op).json()}
    assert comp["ISB-007"]["result_status"] != "not_applicable"


def test_manual_type_is_not_overridden(client, auth):
    from tests.factories import make_story_docx

    op = auth("intake_operator")
    qc = quick(client, op, make_story_docx(child=True), manuscript_type="nonfiksi_umum")
    run = latest_run(client, op, qc["version"]["id"])
    c = classification(client, op, run["id"])
    assert c["applied"] is False and c["manual_type"] == "nonfiksi_umum" and c["jenis_terbitan"]["value"] == "buku_anak"
    assert client.get(f"{API}/manuscripts/{qc['manuscript_id']}", headers=op).json()["manuscript_type"] == "nonfiksi_umum"


def test_min_pages_from_pdf_blocks_but_a4_draft_does_not(client, auth):
    from tests.factories import make_pdf

    op = auth("intake_operator")
    pdf = latest_run(client, op, quick(client, op, make_pdf(), name="buku.pdf", mime="application/pdf", manuscript_type="nonfiksi_umum")["version"]["id"])
    comp = {c["rule_code"]: c for c in client.get(f"{API}/screenings/{pdf['id']}/compliance", headers=op).json()}
    assert comp["ISB-009"]["result_status"] == "fail" and "49" in comp["ISB-009"]["explanation"]
    docx_run = latest_run(client, op, quick(client, op, make_book_docx(), manuscript_type="nonfiksi_umum")["version"]["id"])
    comp = {c["rule_code"]: c for c in client.get(f"{API}/screenings/{docx_run['id']}/compliance", headers=op).json()}
    assert comp["ISB-009"]["observed"] == "unknown" and docx_run["isbn_status"] == "PASS"



# --- Terbitan untuk umum (ISB-014) ---------------------------------------------


def isbn_of(client, headers, qc):
    run = latest_run(client, headers, qc["version"]["id"])
    comp = {c["rule_code"]: c for c in client.get(f"{API}/screenings/{run['id']}/compliance", headers=headers).json()}
    return run, comp


def test_limited_print_statement_is_not_eligible(client, auth):
    import io

    import docx as _docx

    op = auth("intake_operator")
    d = _docx.Document(io.BytesIO(make_book_docx()))
    for p in d.paragraphs:
        if p.text.startswith("Hak cipta"):
            p.insert_paragraph_before("Buku ini dicetak terbatas untuk kalangan sendiri dan tidak diperjualbelikan.")
            break
    buf = io.BytesIO()
    d.save(buf)
    run, comp = isbn_of(client, op, quick(client, op, buf.getvalue(), manuscript_type="nonfiksi_umum"))
    assert run["isbn_status"] == "FAIL" and comp["ISB-014"]["result_status"] == "fail"
    assert "terbatas" in comp["ISB-014"]["explanation"]
    f = next(x for x in findings(client, op, run["id"]) if x["rule_code"] == "ISB-014")
    assert f["paragraph_id"] and "dicetak terbatas" in f["evidence_excerpt"]


def test_narrative_word_terbatas_is_not_a_limited_statement(client, auth):
    import io

    import docx as _docx

    op = auth("intake_operator")
    d = _docx.Document(io.BytesIO(make_book_docx()))
    d.add_paragraph("Keluarga kami hidup dengan penghasilan terbatas dan bekal terbatas.")
    buf = io.BytesIO()
    d.save(buf)
    run, comp = isbn_of(client, op, quick(client, op, buf.getvalue(), manuscript_type="nonfiksi_umum"))
    assert comp["ISB-014"]["result_status"] == "pass" and run["isbn_status"] == "PASS"


def test_declared_distribution_rules(client, auth):
    op = auth("intake_operator")
    run, comp = isbn_of(client, op, quick(client, op, make_book_docx(), manuscript_type="nonfiksi_umum", distribution="terbatas"))
    assert run["isbn_status"] == "FAIL"
    # Biografi/fiksi wajib dijual untuk umum (Juknis 4.5.1.3 butir 6–9): gratis tidak cukup.
    run, comp = isbn_of(client, op, quick(client, op, make_book_docx(), manuscript_type="biografi", distribution="gratis_umum"))
    assert run["isbn_status"] == "FAIL" and "4.5.1.3" in comp["ISB-014"]["explanation"]
    run, _ = isbn_of(client, op, quick(client, op, make_book_docx(), manuscript_type="buku_ajar", distribution="gratis_umum"))
    assert run["isbn_status"] == "PASS"
    # Belum dikonfirmasi → perlu konfirmasi, bukan lulus.
    qc = quick(client, op, make_book_docx(), manuscript_type="biografi", distribution=None)
    run, comp = isbn_of(client, op, qc)
    assert run["isbn_status"] == "NEEDS_VERIFICATION" and comp["ISB-014"]["observed"] == "unknown"
    # Konfirmasi lewat metadata lalu periksa ulang → layak.
    ms = client.get(f"{API}/manuscripts/{qc['manuscript_id']}", headers=op).json()
    client.patch(f"{API}/manuscripts/{qc['manuscript_id']}", json={"metadata": {**ms["metadata_json"], "distribution": "dijual_umum"}}, headers=op)
    client.post(f"{API}/screenings/{run['id']}/rerun", headers=op)
    assert latest_run(client, op, qc["version"]["id"])["isbn_status"] == "PASS"



def _family_memoir(price: bool = False) -> bytes:
    import io

    import docx as _docx

    d = _docx.Document(io.BytesIO(make_book_docx(title="Perjalanan Hidup Rina Pratiwi")))
    copyright_p = next(p for p in d.paragraphs if p.text.startswith("Hak cipta"))
    for line in ["Editor :", "Budi Pratiwi, S.Pd.", "Sari Pratiwi"] + (["Harga: Rp85.000"] if price else []):
        copyright_p.insert_paragraph_before(line)
    paras = d.paragraphs
    i = next(i for i, p in enumerate(paras) if p.text == "KATA PENGANTAR")
    paras[i + 1].insert_paragraph_before("Buku ini merupakan autobiografi yang mengisahkan perjalanan hidup penulis.")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def test_family_memoir_rejected_as_limited(client, auth):
    """Autobiografi orang pribadi yang disusun keluarga ditolak Perpusnas sebagai terbitan kalangan sendiri, walau dinyatakan dijual."""
    op = auth("intake_operator")
    run, comp = isbn_of(client, op, quick(client, op, _family_memoir(), manuscript_type="biografi", distribution="dijual_umum"))
    assert run["isbn_status"] == "FAIL" and comp["ISB-014"]["result_status"] == "fail"
    assert "kalangan sendiri" in comp["ISB-014"]["explanation"]
    f = next(x for x in findings(client, op, run["id"]) if x["rule_code"] == "ISB-014")
    assert "Budi Pratiwi" in f["evidence_excerpt"]
    # Ada bukti komersial (harga di halaman awal) → mengikuti rencana peredaran yang dinyatakan.
    run, _ = isbn_of(client, op, quick(client, op, _family_memoir(price=True), manuscript_type="biografi", distribution="dijual_umum"))
    assert run["isbn_status"] == "PASS"
    # Biografi tanpa tim keluarga tetap lulus bila dijual umum.
    run, _ = isbn_of(client, op, quick(client, op, make_book_docx(title="Perjalanan Hidup Rina Pratiwi"), manuscript_type="biografi"))
    assert run["isbn_status"] == "PASS"


def test_classifies_biography():
    import io

    import docx as _docx

    from app.reader import extract
    from app.screening.classify import classify
    from app.screening.docview import DocView, ManuscriptInfo

    d = _docx.Document(io.BytesIO(make_book_docx()))
    for p in d.paragraphs:
        if p.text == "KATA PENGANTAR":
            p.insert_paragraph_before("Perjalanan Hidup Seorang Anak Desa")
        if p.text.startswith("Pagi itu Kona"):
            p.text = "Buku ini merupakan autobiografi yang menceritakan perjalanan hidup saya sejak kecil."
    buf = io.BytesIO()
    d.save(buf)
    r = extract("docx", buf.getvalue())
    doc = DocView.from_document(r.document, ManuscriptInfo("x", "lainnya", "id", [], {}, False), r.quality.status)
    c = classify(doc, None, "docx")
    assert c["jenis_terbitan"]["value"] == "biografi" and c["peredaran"]["value"] == "Belum dikonfirmasi"


def _reference_book() -> bytes:
    """Buku satu bidang ilmu: daftar isi diketik ber-style Heading, subbab bab terakhir keliru memakai Heading 1."""
    import io

    import docx as _docx

    d = _docx.Document()
    d.add_paragraph("Pengantar Ilmu Ekonomi")
    d.add_paragraph("Penulis: Rina Pratiwi")
    d.add_paragraph("Penerbit Dimensi Antara Pustaka")
    d.add_paragraph("Hak cipta © 2026 pada penulis. Diterbitkan oleh PT Dimensi Antara Pustaka.")
    d.add_heading("KATA PENGANTAR", level=1)
    d.add_paragraph("Buku ini disusun untuk membantu mahasiswa memahami konsep dan prinsip dasar ilmu ekonomi.")
    d.add_heading("DAFTAR ISI", level=1)
    for n in range(1, 7):
        d.add_heading(f"BAB {n} TEORI EKONOMI BAGIAN {n} {n * 10}", level=1)
    for n in range(1, 7):
        d.add_heading(f"BAB {n} TEORI EKONOMI BAGIAN {n}", level=1)
        if n == 6:
            d.add_heading(f"{n}.1 Subbab yang keliru ber-style Heading 1", level=1)
        for _ in range(3):
            d.add_paragraph("Permintaan pasar dipengaruhi harga, pendapatan, dan preferensi konsumen dalam jangka pendek maupun panjang.")
    d.add_heading("DAFTAR PUSTAKA", level=1)
    for i in range(6):
        d.add_paragraph(f"Penulis{i}, A. (20{10 + i}). Judul Buku Ekonomi {i}. Jakarta: Penerbit Ilmu.")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def test_reference_book_and_heading_toc():
    from app.reader import extract
    from app.screening.classify import classify
    from app.screening.docview import DocView, ManuscriptInfo

    r = extract("docx", _reference_book())
    doc = DocView.from_document(r.document, ManuscriptInfo("x", "lainnya", "id", [], {}, False), r.quality.status)
    # Entri daftar isi ber-style Heading tidak dihitung sebagai bab; subbab "6.1" tetap bagian dari BAB 6.
    assert len(doc.chapters) == 6 and doc.chapters[0].title.startswith("TEORI") and doc.chapters[-1].words > 20
    assert classify(doc, None, "docx")["jenis_terbitan"]["value"] == "referensi"


def test_non_commercial_type_needs_no_distribution(client, auth):
    """Buku ajar/referensi diberi ISBN bila dapat diakses umum, berbayar maupun gratis: tanpa pernyataan terbatas → lulus."""
    op = auth("intake_operator")
    run, comp = isbn_of(client, op, quick(client, op, make_book_docx(), manuscript_type="referensi", distribution=None))
    assert comp["ISB-014"]["result_status"] == "pass" and run["isbn_status"] == "PASS"
