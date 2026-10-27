from app.reader import ExtractionError, extract
from app.reader.patterns import classify_heading
from tests.factories import make_docx, make_pdf, make_scanned_pdf, make_txt


def headings(result):
    return [(b.level, b.text) for b in result.document.blocks if b.type == "heading"]


def test_docx_structure_and_locators():
    r = extract("docx", make_docx())
    doc = r.document
    hs = headings(r)
    assert (1, "BAB I PENDAHULUAN") in hs, "BAB I + judul baris berikut digabung"
    assert (2, "1.1 Latar Belakang") in hs
    assert (1, "BAB II KAJIAN PUSTAKA") in hs
    types = {b.type for b in doc.blocks}
    assert {"heading", "paragraph", "list_item", "quote", "table"} <= types
    para = next(b for b in doc.blocks if b.section_path and b.section_path.endswith("1.1 Latar Belakang") and b.type == "paragraph")
    assert para.section_path == "BAB I PENDAHULUAN > 1.1 Latar Belakang"
    assert doc.full_text[para.char_start:para.char_end] == para.text
    assert para.page_number is None and para.locator["page_estimate"] == 1
    kajian = next(b for b in doc.blocks if b.text == "BAB II KAJIAN PUSTAKA")
    assert kajian.locator["page_estimate"] == 2
    assert doc.metadata["title"] == "Judul Uji"
    assert r.quality.status == "EXTRACTED"
    assert any(w.code == "PAGE_NUMBERS_ESTIMATED" for w in doc.warnings)


def test_docx_tracked_changes_and_footnotes():
    r = extract("docx", make_docx(tracked=True, footnote=True))
    text = r.document.full_text
    assert "TEKS SISIPAN" in text
    assert "TEKS TERHAPUS" not in text
    fns = [b for b in r.document.blocks if b.type == "footnote"]
    assert len(fns) == 1 and fns[0].text == "Catatan kaki uji tentang sumber data."
    assert fns[0].section_path.endswith("1.1 Latar Belakang")
    assert any(w.code == "TRACKED_CHANGES" for w in r.document.warnings)


def test_docx_pattern_heading_without_styles():
    r = extract("docx", make_docx(pattern_heading=True))
    assert (1, "BAB I PENDAHULUAN") in headings(r)
    assert any(w.code == "HEADINGS_FROM_PATTERN" for w in r.document.warnings)


def test_pdf_with_outline():
    r = extract("pdf", make_pdf())
    doc = r.document
    hs = headings(r)
    assert (1, "BAB I BAGIAN 1") in hs
    assert (2, "1.1 Subbagian Halaman 1") in hs
    assert doc.page_numbers_reliable
    assert all(b.page_number for b in doc.blocks)
    assert "Header Berjalan" not in doc.full_text, "header berulang dibuang"
    assert doc.pages[0].label == "1"
    assert r.quality.status == "EXTRACTED"


def test_pdf_without_outline_uses_font_and_pattern():
    r = extract("pdf", make_pdf(with_toc=False))
    hs = headings(r)
    assert any(t.startswith("BAB I") and lvl == 1 for lvl, t in hs)
    assert any(t.startswith("1.1 Subbagian") for _, t in hs)


def test_scanned_pdf_is_incomplete_not_scored():
    r = extract("pdf", make_scanned_pdf())
    assert r.quality.status == "INCOMPLETE"
    assert "SCANNED_PAGES" in r.quality.reasons and "LOW_TEXT" in r.quality.reasons
    codes = {w.code for w in r.document.warnings}
    assert {"SCANNED_PAGES", "OCR_NOT_AVAILABLE"} <= codes


def test_corrupt_files_raise_structured_error():
    for fmt, data in (("pdf", b"%PDF-1.7 rusak"), ("docx", b"PK\x03\x04bukan zip")):
        try:
            extract(fmt, data)
        except ExtractionError as exc:
            assert exc.code in ("PDF_UNREADABLE", "UNREADABLE")
        else:
            raise AssertionError(f"{fmt} rusak seharusnya gagal")


def test_txt_reader():
    r = extract("txt", make_txt())
    hs = headings(r)
    assert (1, "BAB I PENDAHULUAN") in hs and (2, "1.1 Latar Belakang") in hs


def test_extraction_is_deterministic():
    data = make_docx(footnote=True)
    a, b = extract("docx", data), extract("docx", data)
    assert a.document.to_summary() == b.document.to_summary()
    assert a.document.full_text == b.document.full_text
    assert a.quality.score == b.quality.score


def test_heading_patterns():
    assert classify_heading("BAB IV Hasil dan Pembahasan") == 1
    assert classify_heading("2.3.1 Metode Sampling") == 3
    assert classify_heading("Daftar Pustaka") == 1
    assert classify_heading("2.3 Ini adalah kalimat biasa yang selesai.") is None
    assert classify_heading("Paragraf biasa tanpa pola heading") is None


def _with_rendered_breaks(data: bytes) -> bytes:
    """Tiru dokumen yang disimpan Word: page break manual + lastRenderedPageBreak di awal paragraf berikutnya."""
    import io
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "word/document.xml":
                text = content.decode()
                i = text.find("BAB II KAJIAN PUSTAKA")
                run = text.rfind("<w:r>", 0, i)
                text = text[:run] + "<w:r><w:lastRenderedPageBreak/></w:r>" + text[run:]
                content = text.encode()
            elif item.filename == "docProps/app.xml":
                import re

                xml = re.sub(r"<Pages>\d+</Pages>", "", content.decode())
                content = xml.replace("</Properties>", "<Pages>2</Pages></Properties>").encode()
            dst.writestr(item, content)
    return out.getvalue()


def test_docx_rendered_breaks_not_double_counted_and_chapters():
    r = extract("docx", _with_rendered_breaks(make_docx()))
    doc = r.document
    kajian = next(b for b in doc.blocks if b.text == "BAB II KAJIAN PUSTAKA")
    assert kajian.locator["page_estimate"] == 2, "page break manual + render break = satu halaman baru"
    s = doc.to_summary()
    assert s["page_count"] == 2 and s["page_count_source"] == "word_app_properties"
    assert [(c["number"], c["title"]) for c in s["chapters"]] == [(1, "PENDAHULUAN"), (2, "KAJIAN PUSTAKA")]


def test_heading_line_breaks_collapsed():
    import docx as _docx
    import io

    d = _docx.Document()
    h = d.add_heading("BAB 1", level=1)
    h.runs[0].add_break()
    h.add_run("DI DESA KECIL")
    d.add_paragraph("Isi bab. " * 100)
    buf = io.BytesIO()
    d.save(buf)
    r = extract("docx", buf.getvalue())
    assert r.document.to_summary()["chapters"][0]["title"] == "DI DESA KECIL"
    assert any(b.text == "BAB 1 DI DESA KECIL" for b in r.document.blocks)
